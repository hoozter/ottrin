"""Generate owned Office fidelity fixtures (python-docx/openpyxl/python-pptx/Pillow).

These are test-data authoring dependencies, never application dependencies.
Keep fixed metadata and ZIP order/timestamps so reruns replace the same fixtures.
"""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import zipfile
import re
from PIL import Image
from docx import Document
from docx.shared import Inches as DocInches, Pt, RGBColor as DocRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.page import PageMargins
from pptx import Presentation
from pptx.util import Inches, Pt as SlidePt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent
STAMP = datetime(2026, 1, 1)
image = Image.new('RGB', (120, 60), '#E87524')
for x in range(60, 120):
    for y in range(60):
        image.putpixel((x, y), (0, 180, 220))
bitmap = BytesIO()
image.save(bitmap, format='PNG')

def save(package, name):
    stream = BytesIO()
    package.save(stream)
    with zipfile.ZipFile(stream) as original, zipfile.ZipFile(ROOT / name, 'w', zipfile.ZIP_DEFLATED) as out:
        for entry in sorted(original.namelist()):
            info = zipfile.ZipInfo(entry, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            data = original.read(entry)
            if entry == 'docProps/core.xml':
                data = re.sub(
                    rb'(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)',
                    rb'\g<1>2026-01-01T00:00:00Z\2', data,
                )
            out.writestr(info, data)

word = Document()
word.core_properties.created = word.core_properties.modified = STAMP
section = word.sections[0]
section.page_width, section.page_height = DocInches(8.5), DocInches(11)
section.left_margin = section.right_margin = DocInches(1)
normal = word.styles['Normal'].font
normal.name, normal.size = 'Arial', Pt(11)
title = word.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('OTTRIN DOCUMENT FIDELITY')
run.font.name, run.font.size, run.bold = 'Arial', Pt(20), True
run.font.color.rgb = DocRGB.from_string('245FA5')
p = word.add_paragraph()
p.add_run('Bold sample').bold = True
p.add_run(' and ')
p.add_run('italic sample').italic = True
word.add_picture(BytesIO(bitmap.getvalue()), width=DocInches(2))
table = word.add_table(rows=3, cols=2)
table.style = 'Table Grid'
for row, values in zip(table.rows, [('Feature', 'Expected layout'), ('Image', 'Orange and cyan halves'), ('Table', 'Two bordered columns')]):
    for cell, value in zip(row.cells, values):
        cell.text = value
word.add_page_break()
p = word.add_paragraph('SECOND PAGE RIGHT ALIGNED')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
save(word, 'office-fidelity.docx')

book = Workbook()
book.properties.created = book.properties.modified = STAMP
sheet = book.active
sheet.title = 'Budget'
sheet.merge_cells('A1:D1')
sheet['A1'] = 'OTTRIN SHEET FIDELITY'
sheet['A1'].font = Font(name='Arial', size=16, bold=True, color='FFFFFF')
sheet['A1'].fill = PatternFill('solid', fgColor='245FA5')
sheet['A1'].alignment = Alignment(horizontal='center')
sheet.row_dimensions[1].height = 30
for row in [('Item', 'Quantity', 'Price', 'Amount'), ('Alpha', 1, 1250, '=B3*C3'), ('Beta', 2, 1375, '=B4*C4'), ('Total', None, None, '=SUM(D3:D4)')]:
    sheet.append(row)
line = Side(style='thin', color='404040')
for row in sheet.iter_rows(min_row=2, max_row=5, max_col=4):
    for cell in row:
        cell.font = Font(name='Arial', size=11, bold=cell.row in (2, 5))
        cell.border = Border(left=line, right=line, top=line, bottom=line)
        if cell.column in (3, 4) and cell.row > 2:
            cell.number_format = '"$"#,##0.00'
sheet.merge_cells('A7:D7')
sheet['A7'] = 'Currency totals and calculated formulas'
for col in 'ABCDEFGH':
    sheet.column_dimensions[col].width = 14
chart = BarChart()
chart.title = 'Amounts by item'
chart.add_data(Reference(sheet, min_col=4, min_row=2, max_row=4), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=4))
chart.series[0].graphicalProperties.solidFill = '2C9C69'
chart.width, chart.height = 9, 7
sheet.add_chart(chart, 'E2')
sheet.print_area = 'A1:I18'
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.page_setup.orientation = 'landscape'
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
sheet.page_margins = PageMargins(left=.25, right=.25, top=.25, bottom=.25, header=0, footer=0)
save(book, 'office-fidelity.xlsx')

slides = Presentation()
slides.core_properties.created = slides.core_properties.modified = STAMP
slides.slide_width, slides.slide_height = Inches(10), Inches(5.625)
for number in (1, 2):
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(9), Inches(.6)).text_frame
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f'OTTRIN SLIDE {number} FIDELITY'
    r.font.name, r.font.size, r.font.bold = 'Arial', SlidePt(24), True
    r.font.color.rgb = RGBColor.from_string('245FA5')
    if number == 1:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.7), Inches(1.5), Inches(4), Inches(2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string('2C9C69')
        shape.line.color.rgb = RGBColor.from_string('165C38')
        shape.text = 'Green shape with text'
        slide.shapes.add_picture(BytesIO(bitmap.getvalue()), Inches(6), Inches(1.5), width=Inches(2.5))
    else:
        frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame
        frame.text = 'Second slide stays separate and ordered'
save(slides, 'office-fidelity.pptx')
